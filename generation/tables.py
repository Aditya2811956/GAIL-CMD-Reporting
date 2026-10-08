"""
DOCX table manipulation utilities.

Handles:
- Style snapshotting and restoration
- Table border formatting
- Markdown to DOCX table conversion
"""

import copy
import logging

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)


def snapshot_table_styles(table: Table) -> List[List[Dict[str, Any]]]:
    """
    Capture style information for every cell in a table.

    Returns a 2-D list [row][col] of deep-copied style blobs.

    For each cell we capture:
        - tcPr (borders, shading, width)
        - pPr (paragraph properties, alignment)
        - rPr (run properties: font, size, bold)
    """

    snap = []

    for row in table.rows:
        snap_row = []

        for cell in row.cells:
            cell_style = {
                "tcPr": copy.deepcopy(cell._tc.tcPr),
                "pPr": (
                    copy.deepcopy(cell.paragraphs[0]._p.pPr)
                    if cell.paragraphs
                    else None
                ),
                "rPr": (
                    copy.deepcopy(
                        cell.paragraphs[0].runs[0]._r.rPr
                    )
                    if cell.paragraphs
                    and cell.paragraphs[0].runs
                    else None
                ),
            }

            snap_row.append(cell_style)

        snap.append(snap_row)

    return snap


def restore_table_styles(
    table: Table,
    snap: List[List[Dict[str, Any]]]
) -> None:
    """
    Apply cached styles to the whole table.

    Handles:
    - Cells that existed before -> original style
    - Extra columns -> style copied from last old column
    - Extra rows -> style copied from last old row
    """

    orig_rows = len(snap)
    orig_cols = len(snap[0]) if snap else 0

    if orig_rows == 0 or orig_cols == 0:
        return

    for r, row in enumerate(table.rows):
        row_src = snap[min(r, orig_rows - 1)]

        for c, cell in enumerate(row.cells):
            src_style = row_src[min(c, orig_cols - 1)]

            # tcPr (cell properties)
            dst_tcPr = cell._tc.get_or_add_tcPr()
            dst_tcPr.clear()

            if src_style["tcPr"] is not None:
                for child in src_style["tcPr"]:
                    dst_tcPr.append(copy.deepcopy(child))

            # pPr (paragraph properties)
            if cell.paragraphs:
                dst_pPr = cell.paragraphs[0]._p.get_or_add_pPr()
                dst_pPr.clear()

                if src_style["pPr"] is not None:
                    for child in src_style["pPr"]:
                        dst_pPr.append(copy.deepcopy(child))

            # rPr (run properties: font/bold/size)
            if cell.paragraphs[0].runs:
                dst_rPr = (
                    cell.paragraphs[0]
                    .runs[0]
                    ._r
                    .get_or_add_rPr()
                )

                dst_rPr.clear()

                if src_style["rPr"] is not None:
                    for child in src_style["rPr"]:
                        dst_rPr.append(copy.deepcopy(child))


def force_black_grid(
    table: Table,
    size: int = 4,
    color: str = "000000"
) -> None:
    """
    Apply a black grid (outer + inner borders) to a Word table.
    """

    tbl = table._tbl

    # Get or create tblPr
    tbl_pr = tbl.find(qn("w:tblPr"))

    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    # Get or create tblBorders
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))

    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)

    # Set all edges
    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        edge_elm = tbl_borders.find(qn(f"w:{edge}"))

        if edge_elm is None:
            edge_elm = OxmlElement(f"w:{edge}")
            tbl_borders.append(edge_elm)

        edge_elm.set(qn("w:val"), "single")
        edge_elm.set(qn("w:sz"), str(size))
        edge_elm.set(qn("w:color"), color)
        edge_elm.set(qn("w:space"), "0")


def markdown_to_docx_table(md_text: str) -> Document:
    """
    Convert a markdown table to a python-docx Table.

    Returns:
        Document containing the table.
        Access via doc.tables[0].
    """

    rows = [
        [cell.strip() for cell in row.split("|")[1:-1]]
        for row in md_text.strip().splitlines()
        if row.startswith("|") and "---" not in row
    ]

    frag_doc = Document()

    if not rows:
        return frag_doc

    tbl = frag_doc.add_table(
        rows=len(rows),
        cols=len(rows[0])
    )

    tbl.style = "Table Grid"

    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            tbl.rows[r].cells[c].text = cell

    return frag_doc


def set_table_column_widths(
    table: Table,
    widths: List[float],
) -> None:
    """
    Set explicit column widths for a DOCX table.

    widths are specified in inches.
    """

    from docx.shared import Inches

    table.autofit = False

    logger.warning("===== LNG WIDTH FUNCTION CALLED =====")

    tbl = table._tbl

    # Remove existing table grid
    old_grid = tbl.find(qn("w:tblGrid"))

    if old_grid is not None:
        tbl.remove(old_grid)

    # Create a new table grid
    tbl_grid = OxmlElement("w:tblGrid")

    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(
            qn("w:w"),
            str(int(width * 1440))
        )
        tbl_grid.append(grid_col)

    # Insert grid immediately after tblPr
    tbl_pr = tbl.find(qn("w:tblPr"))

    if tbl_pr is not None:
        tbl_pr.addnext(tbl_grid)
    else:
        tbl.insert(0, tbl_grid)

    # Set individual cell widths
    for row in table.rows:
        for idx, width in enumerate(widths):

            if idx >= len(row.cells):
                continue

            cell = row.cells[idx]

            cell.width = Inches(width)

            tc_pr = cell._tc.get_or_add_tcPr()

            tc_w = tc_pr.find(qn("w:tcW"))

            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)

            tc_w.set(
                qn("w:w"),
                str(int(width * 1440))
            )

            tc_w.set(
                qn("w:type"),
                "dxa"
            )


def update_table_from_markdown(
    doc: Document,
    table_index: int,
    markdown_table: str,
    preserve_styles: bool = True
) -> None:
    """
    Replace an existing DOCX table with a table generated
    from Markdown.

    If preserve_styles=True, the original table's formatting
    is restored after replacing the rows.
    """

    if table_index >= len(doc.tables):
        raise IndexError(
            f"Table index {table_index} out of range. "
            f"Document has {len(doc.tables)} tables."
        )

    target_table = doc.tables[table_index]

    frag_doc = markdown_to_docx_table(markdown_table)

    if not frag_doc.tables:
        raise ValueError(
            "Markdown table could not be converted to a DOCX table."
        )

    new_table = frag_doc.tables[0]

    if preserve_styles:

        style_snapshot = snapshot_table_styles(
            target_table
        )

        target_table._tbl.clear()

        for r in new_table.rows:
            target_table._tbl.append(
                copy.deepcopy(r._tr)
            )

        restore_table_styles(
            target_table,
            style_snapshot
        )

        force_black_grid(
            target_table
        )

    else:

        target_table._tbl.clear()

        for r in new_table.rows:
            target_table._tbl.append(
                copy.deepcopy(r._tr)
            )


def get_table_markdown(table: Table) -> str:
    """
    Convert a python-docx table to markdown format.
    """

    from tabulate import tabulate

    data = [
        [cell.text.strip() for cell in row.cells]
        for row in table.rows
    ]

    return tabulate(
        data,
        headers="firstrow",
        tablefmt="github"
    )


def add_update_note_after_table(
    doc: Document,
    table: Table,
    note_text: Optional[str] = None
) -> None:
    """
    Insert a small, italic 'Updated ...' note right under a table.
    """

    from datetime import datetime
    from docx.shared import Pt, RGBColor
    from docx.text.paragraph import Paragraph

    if note_text is None:
        note_text = (
            f"Updated on {datetime.now():%d %b %Y}"
        )

    # Insert paragraph after table
    new_p_elm = OxmlElement("w:p")

    table._tbl.addnext(new_p_elm)

    p = Paragraph(
        new_p_elm,
        doc
    )

    run = p.add_run(note_text)

    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(
        0x59,
        0x59,
        0x59
    )